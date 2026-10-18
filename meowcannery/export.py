"""Exameow 1.5.0 原生 15 列导入契约，文本保真、原子写入。"""
import os
import re
import tempfile
from pathlib import Path
from zipfile import ZipFile

from .validation import validate_bank

TARGET_VERSION = "1.5.0"
COMPATIBLE_VERSIONS = ("1.4.4", "1.4.5", "1.5.0", "1.6.0", "1.6.1")
COMPATIBLE_RANGE = "1.4.4–1.6.1"
HEADERS = ["题干（必填）", "题型 （必填）"] + [f"选项 {chr(65+i)}" for i in range(8)] + [
    "正确答案\n（必填）", "解析\n（勿删）", "学科", "章节\n（勿删）", "难度"]


def rows_for(questions, subject=""):
    return [[q["stem"], q["type"], *q["options"], *([""] * (8 - len(q["options"]))),
             q["answer"], q.get("explanation", ""), q.get("subject", subject), q["chapter"],
             q.get("difficulty", "")] for q in questions]


def read_xlsx(path):
    from openpyxl import load_workbook
    with Path(path).open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=False)
        try:
            rows = workbook.worksheets[0].iter_rows(values_only=True)
            header = next(rows, ())
            if list(header) != HEADERS:
                raise ValueError("表头不符合 Exameow 1.5.0 的原生 15 列格式。")
            questions = []
            for row in rows:
                values = ["" if v is None else str(v).strip() for v in row]
                if not any(values):
                    continue
                values += [""] * (15 - len(values))
                if any(not values[i] and any(values[j] for j in range(i+1, 10)) for i in range(2, 10)):
                    raise ValueError("选项列中存在空洞，导入后会改变答案对应关系。")
                questions.append(dict(stem=values[0], type=values[1], options=[v for v in values[2:10] if v],
                                      answer=values[10], explanation=values[11], subject=values[12],
                                      chapter=values[13], difficulty=values[14]))
            return questions
        finally:
            workbook.close()


def _preserve_non_bmp(path):
    """避免旧版 XLSX 读取器将非 BMP 数字实体截断为 16 位字符。"""
    def literal_utf8(match):
        value = int(match[1], 16) if match[1] else int(match[2], 10)
        return chr(value).encode("utf-8") if 0x10000 <= value <= 0x10FFFF else match[0]

    entity = re.compile(rb"&#(?:[xX]([0-9a-fA-F]+)|([0-9]+));")
    rewritten = None
    try:
        with ZipFile(path) as source:
            changed = {}
            for entry in source.infolist():
                if (entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml")
                        or entry.filename == "xl/sharedStrings.xml"):
                    original = source.read(entry)
                    updated = entity.sub(literal_utf8, original)
                    if updated != original:
                        changed[entry.filename] = updated
            if not changed:
                return
            fd, filename = tempfile.mkstemp(prefix=path.stem + ".utf8.", suffix=".xlsx", dir=path.parent)
            os.close(fd)
            rewritten = Path(filename)
            with ZipFile(rewritten, "w") as target:
                target.comment = source.comment
                for entry in source.infolist():
                    data = changed.get(entry.filename)
                    target.writestr(entry, source.read(entry) if data is None else data)
        # Windows 下关闭源 ZIP 后才能替换，且始终只改写导出过程的临时文件。
        rewritten.replace(path)
    finally:
        if rewritten is not None:
            rewritten.unlink(missing_ok=True)


def write_xlsx(questions, path, subject=""):
    # 项目运行时沿用原有 openpyxl，无需 Node 或 Codex 的专有运行环境。
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    validate_bank(questions)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    expected = rows_for(questions, subject)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "题库"
    for row in [HEADERS, *expected]:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # 化学等式以 = 开头时也必须作为原文保存。
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.font = Font(name="Microsoft YaHei", size=10)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    widths = [60, 12] + [23] * 8 + [45, 55, 12, 35, 10]
    for i, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="263B50")
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    sheet.row_dimensions[1].height = 32
    fd, filename = tempfile.mkstemp(prefix=path.stem + ".", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    temporary = Path(filename)
    try:
        workbook.save(temporary)
        _preserve_non_bmp(temporary)
        actual = rows_for(read_xlsx(temporary))
        if actual != [[v.strip() for v in row] for row in expected]:
            raise ValueError("Excel 回读与题库内容不一致，未覆盖输出。")
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return path
